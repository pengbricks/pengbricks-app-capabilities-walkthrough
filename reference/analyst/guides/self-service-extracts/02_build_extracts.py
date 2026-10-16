# Databricks notebook source
# MAGIC %pip install -q -r ../../requirements.txt

# COMMAND ----------

dbutils.library.restartPython()

# COMMAND ----------

# MAGIC %md
# MAGIC # 2. Build the extracts
# MAGIC
# MAGIC Writes a filtered extract as a formatted Excel workbook and a pipe-delimited file for external partners. Files are named by date, so a rerun on the same day replaces them. The table behind the dashboard is overwritten too.

# COMMAND ----------

import sys, os
sys.path.insert(0, os.path.abspath("../../src"))
from guide_params import bundle_root, load_params
params = load_params(dbutils, "capability_guides_extracts", ['catalog', 'schema', 'use_case', 'warehouse_id', 'notification_email', 'dashboard_id', 'run_id'], defaults={'run_id': 'manual'})
catalog, schema = params["catalog"], params["schema"]
fq = f"{catalog}.{schema}"
ROOT = bundle_root()
SPEC_FILE = ROOT / "use_cases" / f"{params['use_case']}.yml"

# COMMAND ----------

# MAGIC %md
# MAGIC ## Filter the orders
# MAGIC
# MAGIC The extract table also backs the dashboard, so the dashboard and the files share one filter, defined in the spec.

# COMMAND ----------

import shutil, tempfile
from datetime import date
from pyspark.sql import functions as F
from analyst_common.helpers import read_delimited, write_delimited
from analyst_common.spec import load_spec

flt = load_spec(SPEC_FILE)["extract"]
orders = spark.table(f"{fq}.analyst_orders").alias("o")
customers = spark.table(f"{fq}.analyst_customers").alias("c")
extract = (
    orders.join(customers, "customer_id")
    .where(
        (F.col("region") == flt["region"])
        & (F.col("status") == flt["status"])
        & (F.col("order_ts") >= F.lit(flt["begin"]))
        & (F.col("order_ts") < F.date_add(F.lit(flt["end"]).cast("date"), 1))
    )
    .select(
        "order_id", "customer_id", "region", "segment", "product_id",
        F.to_date("order_ts").alias("order_date"), "quantity", "unit_price",
        (F.col("quantity") * F.col("unit_price")).cast("decimal(12,2)").alias("revenue"),
    )
    .orderBy("order_date", "order_id")
)
extract.write.mode("overwrite").option("overwriteSchema", "true").saveAsTable(f"{fq}.extract_orders")
pdf = spark.table(f"{fq}.extract_orders").orderBy("order_date", "order_id").toPandas()
header = list(pdf.columns)
print(f"Extract has {len(pdf)} rows and revenue {pdf['revenue'].sum()}")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Write the Excel workbook
# MAGIC
# MAGIC Volume paths do not support the seeks that zip files need, so the workbook is built locally and copied.

# COMMAND ----------

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

folder = f"/Volumes/{catalog}/{schema}/extracts/partner"
os.makedirs(folder, exist_ok=True)
stamp = date.today().isoformat()
xlsx_path = f"{folder}/orders_extract_{stamp}.xlsx"
pipe_path = f"{folder}/orders_extract_{stamp}.txt"
work = tempfile.mkdtemp()
local_xlsx = os.path.join(work, "extract.xlsx")
local_pipe = os.path.join(work, "extract.txt")

wb = Workbook()
ws = wb.active
ws.title = "Orders"
ws.append(header)
for row in pdf.itertuples(index=False):
    ws.append([float(v) if hasattr(v, "quantize") else v for v in row])
head_fill = PatternFill("solid", fgColor="1B3139")
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.alignment = Alignment(horizontal="center")
money_cols = {header.index("unit_price") + 1, header.index("revenue") + 1}
date_col = header.index("order_date") + 1
for row in ws.iter_rows(min_row=2):
    for cell in row:
        if cell.column in money_cols:
            cell.number_format = "#,##0.00"
        elif cell.column == date_col:
            cell.number_format = "yyyy-mm-dd"
for index, name in enumerate(header, start=1):
    width = max([len(name)] + [len(str(c.value)) for c in ws[get_column_letter(index)][1:200]]) + 2
    ws.column_dimensions[get_column_letter(index)].width = width
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
summary = wb.create_sheet("Summary")
summary.append(["Filter", "Value"])
for key, value in flt.items():
    summary.append([key, str(value)])
summary.append(["rows", len(pdf)])
summary.append(["revenue", float(pdf["revenue"].sum())])
for cell in summary[1]:
    cell.font = Font(bold=True)
summary.column_dimensions["A"].width = 14
summary.column_dimensions["B"].width = 16
wb.save(local_xlsx)
shutil.copyfile(local_xlsx, xlsx_path)

# COMMAND ----------

# MAGIC %md
# MAGIC ## Write the pipe-delimited file and verify both

# COMMAND ----------

from openpyxl import load_workbook

count = write_delimited(local_pipe, header, pdf.astype(object).itertuples(index=False, name=None))
shutil.copyfile(local_pipe, pipe_path)

# Read both files back from the volume and confirm they hold the same rows as the table.
xlsx_rows = load_workbook(xlsx_path, read_only=True)["Orders"].max_row - 1
pipe_rows = len(read_delimited(pipe_path))
if not (xlsx_rows == count == pipe_rows == len(pdf)):
    raise Exception(f"Row counts differ: table {len(pdf)}, xlsx {xlsx_rows}, pipe {pipe_rows}")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Record a manifest
# MAGIC
# MAGIC Overwritten on each run.

# COMMAND ----------

manifest = [
    (os.path.basename(p), p, kind, len(pdf), os.path.getsize(p), params["run_id"])
    for p, kind in ((xlsx_path, "xlsx"), (pipe_path, "pipe-delimited"))
]
spark.createDataFrame(
    manifest, "file_name STRING, path STRING, format STRING, row_count BIGINT, bytes BIGINT, run_id STRING"
).withColumn("recorded_at", F.current_timestamp()).write.mode("overwrite").option(
    "overwriteSchema", "true"
).saveAsTable(f"{fq}.extract_manifest")
display(spark.table(f"{fq}.extract_manifest"))
