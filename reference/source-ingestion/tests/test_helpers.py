from __future__ import annotations

import stat
import sys
from decimal import Decimal
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from source_ingestion import api_client, checksum, datagen, file_writers, sftp_utils  # noqa: E402
from source_ingestion.spec import load_spec  # noqa: E402

SPEC = load_spec(str(Path(__file__).resolve().parents[1] / "use_cases" / "general.yml"))


class FakeResponse:
    def __init__(self, status, body=None, headers=None):
        self.status_code, self._body, self.headers = status, body or {}, headers or {}

    def json(self):
        return self._body


class FakeSession:
    """Serves 10 users in pages and records every call. `script` maps call number to a status override."""

    def __init__(self, total=10, script=None):
        self.total, self.script, self.calls, self.gets = total, script or {}, [], 0

    def post(self, url, json=None, timeout=None):
        self.calls.append(url)
        if url.endswith("/auth/login"):
            return FakeResponse(200, {"accessToken": "a1", "refreshToken": "r1"})
        return FakeResponse(200, {"accessToken": f"a{len(self.calls)}", "refreshToken": "r1"})

    def get(self, url, headers=None, params=None, timeout=None):
        self.gets += 1
        if self.gets in self.script:
            status = self.script[self.gets]
            return FakeResponse(status, headers={"Retry-After": "2"} if status == 429 else {})
        skip, limit = params["skip"], params["limit"]
        users = [{"id": i} for i in range(skip, min(skip + limit, self.total))]
        return FakeResponse(200, {"users": users, "total": self.total})


def run(session, **kwargs):
    tokens = api_client.TokenManager(session, "http://x", "u", "p", refresh_after_seconds=1000)
    sleeps = []
    result = api_client.fetch_all_pages(
        session, tokens, "http://x/users", "users", 4, sleep=sleeps.append, **kwargs
    )
    return result, sleeps


def test_follows_pages_until_total():
    result, _ = run(FakeSession(total=10))
    assert [r["id"] for r in result.records] == list(range(10))
    assert [row["skip"] for row in result.log] == [0, 4, 8]


def test_401_refreshes_token_and_retries_without_backoff():
    session = FakeSession(script={2: 401})
    result, sleeps = run(session)
    assert len(result.records) == 10
    assert result.token_refreshes == 1
    assert sleeps == []
    assert [row["http_status"] for row in result.log][:3] == [200, 401, 200]
    assert [row["attempt"] for row in result.log][:3] == [1, 1, 2]


def test_503_retries_with_exponential_backoff_and_429_honors_retry_after():
    result, sleeps = run(FakeSession(script={1: 503, 2: 503, 3: 429}), backoff_base=0.5)
    assert len(result.records) == 10
    assert sleeps == [0.5, 1.0, 2.0]
    assert result.retries == 3


def test_gives_up_after_max_retries_and_stops_on_other_4xx():
    with pytest.raises(api_client.ApiError, match="after 3 attempts"):
        run(FakeSession(script={i: 500 for i in range(1, 20)}), max_retries=2)
    with pytest.raises(api_client.ApiError, match="HTTP 404"):
        run(FakeSession(script={1: 404}))


def test_injected_invalid_token_triggers_refresh_path():
    class Strict(FakeSession):
        def get(self, url, headers=None, params=None, timeout=None):
            if headers["Authorization"].startswith("Bearer invalid"):
                self.gets += 1
                return FakeResponse(401)
            return super().get(url, headers=headers, params=params)

    result, _ = run(Strict(), inject_invalid_token_on_page=2)
    assert len(result.records) == 10
    assert result.token_refreshes == 1


def test_token_refreshes_proactively_when_old():
    clock = iter(range(0, 1000))
    session = FakeSession()
    tokens = api_client.TokenManager(session, "http://x", "u", "p", refresh_after_seconds=2, clock=lambda: next(clock))
    tokens.header()  # login at t=0..
    tokens.header()
    tokens.header()
    assert tokens.logins == 1 and tokens.refreshes >= 1


def test_backoff_caps():
    assert api_client.backoff_delay(10, 0.5) == 30.0
    assert api_client.backoff_delay(0, 0.5, retry_after=100) == 30.0


def test_delimited_quotes_embedded_delimiters(tmp_path):
    rows = [{"id": 1, "notes": "gift wrap | express", "amount": Decimal("1.50")}, {"id": 2, "notes": 'says "rush"', "amount": Decimal("2")}]
    path = tmp_path / "o.txt"
    file_writers.write_delimited(str(path), rows, delimiter="|")
    text = path.read_text().splitlines()
    assert text[0] == "id|notes|amount"
    assert text[1] == '1|"gift wrap | express"|1.50'
    assert text[2] == '2|"says ""rush"""|2'
    assert file_writers.count_delimited_with_embedded_delimiter(rows, "|") == 1


def test_workbook_has_named_sheets(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "w.xlsx"
    file_writers.write_workbook(str(path), {"README": [{"note": "x"}], "Orders": [{"order_id": 1, "amount": Decimal("2.5")}]})
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["README", "Orders"]
    assert [c.value for c in workbook["Orders"][2]] == [1, 2.5]


def test_xml_counts_match_writer(tmp_path):
    orders = [{"order_id": i, "customer_id": 1, "order_date": "2026-01-01", "discount_pct": 0} for i in range(10)]
    tree = file_writers.build_orders_xml(orders, max_lines=4, coupon_every=4)
    root = tree.getroot()
    expected = file_writers.expected_xml_counts(10, 4, 4)
    assert len(root.findall("order")) == expected["orders"]
    assert len(root.findall("order/lines/line")) == expected["lines"]
    assert len(root.findall("order/coupon")) == expected["coupons"]


def test_nested_json_counts():
    customers = [
        {"customer_id": i, "first_name": "A", "last_name": "B", "email": "e", "city": "c", "country": "US",
         "credit_limit": Decimal("1.00"), "created_date": "2026-01-01"}
        for i in range(30)
    ]
    docs = [file_writers.nest_customer(row, i) for i, row in enumerate(customers)]
    expected = file_writers.expected_json_counts(30)
    assert sum(len(d["contacts"]) for d in docs) == expected["contacts"]
    assert sum("loyalty" in d for d in docs) == expected["loyalty"]


def test_checksum_sql_is_symmetric():
    cols = [checksum.Column("customer_id", "int"), checksum.Column("amount", "decimal"), checksum.Column("order_date", "date")]
    pg = checksum.postgres_checksum_sql(cols, "SELECT * FROM t")
    sp = checksum.spark_checksum_sql(cols, "c.s.t")
    assert "to_char(amount, 'FM9999999990.00')" in pg and "md5(" in pg
    assert "CAST(CAST(amount AS DECIMAL(12,2)) AS STRING)" in sp and "FROM c.s.t" in sp


class Entry:
    def __init__(self, filename, mode, size=0, mtime=0):
        self.filename, self.st_mode, self.st_size, self.st_mtime = filename, mode, size, mtime


class FakeSftp:
    tree = {
        "/": [Entry("a.txt", stat.S_IFREG, 5, 10), Entry("pub", stat.S_IFDIR)],
        "/pub": [Entry("b.png", stat.S_IFREG, 7, 20)],
    }

    def listdir_attr(self, path):
        return self.tree[path]


def test_sftp_walk_plan_and_classify():
    remote = sftp_utils.walk_remote(FakeSftp(), "/")
    assert [f.path for f in remote] == ["/a.txt", "/pub/b.png"]
    seen = {("/a.txt", 5, 10)}
    assert [f.path for f in sftp_utils.plan_transfers(remote, seen)] == ["/pub/b.png"]
    changed = {("/a.txt", 6, 10)}
    assert len(sftp_utils.plan_transfers(remote, changed)) == 2
    assert sftp_utils.local_path("/Volumes/c/s/v/sftp", "/pub/b.png") == "/Volumes/c/s/v/sftp/pub/b.png"
    assert sftp_utils.classify("/x/A.TXT", [".txt"], [".png"]) == "text"
    assert sftp_utils.classify("/x/b.png", [".txt"], [".png"]) == "binary"


def test_entity_frame_passes_the_spec_to_the_generator(monkeypatch):
    class Frame:
        def orderBy(self, name):
            return ("ordered", name)

    calls = {}

    def fake_generate(spark, columns, rows, seed, name, partitions=8, **_):
        calls.update(columns=columns, rows=rows, seed=seed, name=name)
        return Frame()

    monkeypatch.setattr(datagen, "generate", fake_generate)
    entity = SPEC["entities"]["customers"]
    result = datagen.entity_frame(None, SPEC, "customers", rows=7, seed_offset=11)
    assert result == ("ordered", entity["columns"][0]["name"])
    assert calls["columns"] is entity["columns"] and calls["rows"] == 7
    assert calls["seed"] == int(SPEC["seed"]) + 11 and calls["name"] == f"{SPEC['name']}_customers"


def test_flatten_user_keeps_nested_fields_and_tolerates_missing_ones():
    from source_ingestion.api_records import flatten_user

    row = flatten_user({"id": 7, "firstName": "Ada", "address": {"city": "Leeds", "country": "UK"}, "company": {"name": "Acme"}})
    assert row["user_id"] == 7 and row["city"] == "Leeds" and row["company"] == "Acme"
    bare = flatten_user({"id": 8})
    assert bare["city"] is None and bare["company"] is None
