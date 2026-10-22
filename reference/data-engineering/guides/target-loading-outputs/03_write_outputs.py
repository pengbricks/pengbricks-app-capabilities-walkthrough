# Databricks notebook source
# MAGIC %pip install -q -r ../../requirements.txt

# COMMAND ----------

dbutils.library.restartPython()

# COMMAND ----------

# MAGIC %md
# MAGIC # 3. Write the file outputs
# MAGIC
# MAGIC Safe to rerun: the four files keep fixed names and are overwritten, and `target_output_files` is overwritten. Writes the merged target to a volume as CSV, pipe-delimited text, Excel and JSON. Files keep a fixed name, so a rerun overwrites them, and `target_output_files` is overwritten too.

# COMMAND ----------

import sys, os
sys.path.insert(0, os.path.abspath("../../src"))
from guide_params import bundle_root, load_params
params = load_params(dbutils, "capability_guides_target_loading", ['catalog', 'schema', 'use_case', 'volume'], defaults={'volume': 'exports'})
catalog, schema = params["catalog"], params["schema"]
scope = f"{catalog}.{schema}"
ROOT = bundle_root()
USE_CASE_FILE = ROOT / "use_cases" / f"{params['use_case']}.yml"

# COMMAND ----------

# MAGIC %md
# MAGIC ## Read the target

# COMMAND ----------

import shutil, tempfile
from datetime import datetime, timezone
from pyspark.sql import functions as F
from de_common.files import column_widths
from de_common.spec import load_use_case

spec = load_use_case(USE_CASE_FILE)
outputs = spec["accounts"]["outputs"]
out_dir = f"/Volumes/{catalog}/{schema}/{params['volume']}"
base = outputs["base_name"]

df = (
    spark.table(f"{scope}.accounts_target")
    .orderBy("account_id")
    .limit(int(outputs["max_rows"]))
    .withColumn("balance", F.col("balance").cast("double"))
    .withColumn("updated_at", F.date_format("updated_at", "yyyy-MM-dd HH:mm:ss"))
)
pdf = df.toPandas()
headers = list(pdf.columns)

# COMMAND ----------

# MAGIC %md
# MAGIC ## Excel writer
# MAGIC
# MAGIC Bold header, widths fitted to the longest value, frozen header row and a number format on decimal columns.

# COMMAND ----------

def write_xlsx(path: str, headers: list[str], rows: list[list], decimal_columns: set[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "accounts"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="1B3139")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for index, width in enumerate(column_widths(headers, rows), start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for index, name in enumerate(headers, start=1):
        if name in decimal_columns:
            for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    item.number_format = "#,##0.00"
    sheet.freeze_panes = "A2"
    workbook.save(path)

# COMMAND ----------

# MAGIC %md
# MAGIC ## Write the four files

# COMMAND ----------

pdf.to_csv(f"{out_dir}/{base}.csv", index=False)
pdf.to_csv(f"{out_dir}/{base}.txt", index=False, sep="|")
# openpyxl seeks inside the zip, which volume paths do not support, so build locally and copy.
with tempfile.TemporaryDirectory() as scratch:
    local = os.path.join(scratch, f"{base}.xlsx")
    write_xlsx(local, headers, pdf.values.tolist(), {"balance"})
    shutil.copyfile(local, f"{out_dir}/{base}.xlsx")
pdf.to_json(f"{out_dir}/{base}.json", orient="records", indent=2)

# COMMAND ----------

# MAGIC %md
# MAGIC ## Record the files

# COMMAND ----------

run_ts = datetime.now(timezone.utc).isoformat(timespec="seconds")
formats = {"csv": "comma-delimited", "txt": "pipe-delimited", "xlsx": "Excel workbook", "json": "JSON array"}
rows = []
for extension, description in formats.items():
    name = f"{base}.{extension}"
    size = os.path.getsize(f"{out_dir}/{name}")
    rows.append((run_ts, name, description, int(size), len(pdf)))
    print(f"{name:<24} {size:>9,} bytes  {description}")
spark.createDataFrame(
    rows, "run_ts string, file_name string, format string, size_bytes long, row_count long"
).write.mode("overwrite").option("overwriteSchema", "true").saveAsTable(f"{scope}.target_output_files")
display(spark.table(f"{scope}.target_output_files"))
